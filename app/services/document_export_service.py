import io
import re

from docx import Document as DocxDocument
from docx.shared import Pt
from fpdf import FPDF, XPos, YPos


def _strip_think_tags(text: str) -> str:
    """Remove <think>...</think> reasoning blocks some models emit."""
    return re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL).strip()


def build_docx(title: str, messages) -> bytes:
    """Render a conversation transcript as a .docx file, returned as bytes."""

    document = DocxDocument()

    document.add_heading(title or "Conversation", level=1)

    for message in messages:

        role_label = "You" if message.role == "user" else "Assistant"

        heading = document.add_paragraph()
        run = heading.add_run(role_label)
        run.bold = True
        run.font.size = Pt(11)

        content = _strip_think_tags(message.content)
        document.add_paragraph(content)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(title: str, messages) -> bytes:
    """Render a conversation transcript as a PDF file, returned as bytes."""

    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # new_x/new_y explicitly reset the cursor to the left margin after each
    # multi_cell — without this, fpdf2 can leave the cursor wherever the
    # previous cell's text happened to end, and a later multi_cell(0, ...)
    # ("use full remaining width") can compute that remaining width as
    # ~0 and crash with "Not enough horizontal space to render a single
    # character" even for short, ordinary text.
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, _to_latin1(title or "Conversation"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)

    for message in messages:

        role_label = "You" if message.role == "user" else "Assistant"

        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 8, _to_latin1(role_label), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        pdf.set_font("Helvetica", "", 11)
        content = _strip_think_tags(message.content)
        pdf.multi_cell(0, 6, _to_latin1(content), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(3)

    return bytes(pdf.output())


def _to_latin1(text: str) -> str:
    """fpdf2's built-in fonts only support latin-1; swap unsupported chars
    (emoji, curly quotes, etc.) for a safe placeholder instead of crashing."""

    return text.encode("latin-1", errors="replace").decode("latin-1")
