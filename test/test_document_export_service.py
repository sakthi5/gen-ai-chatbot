import io

from docx import Document as DocxDocument

from app.services.document_export_service import build_docx, build_pdf


class FakeMessage:
    def __init__(self, role, content):
        self.role = role
        self.content = content


def test_build_docx_contains_conversation_text():
    messages = [
        FakeMessage("user", "What is RAG?"),
        FakeMessage("assistant", "RAG stands for Retrieval-Augmented Generation."),
    ]

    file_bytes = build_docx("RAG definition", messages)

    document = DocxDocument(io.BytesIO(file_bytes))
    full_text = "\n".join(p.text for p in document.paragraphs)

    assert "RAG definition" in full_text
    assert "What is RAG?" in full_text
    assert "Retrieval-Augmented Generation" in full_text


def test_build_docx_strips_think_tags():
    messages = [
        FakeMessage("assistant", "<think>internal reasoning</think>The answer is 42."),
    ]

    file_bytes = build_docx("Test", messages)
    document = DocxDocument(io.BytesIO(file_bytes))
    full_text = "\n".join(p.text for p in document.paragraphs)

    assert "internal reasoning" not in full_text
    assert "The answer is 42." in full_text


def test_build_pdf_returns_valid_pdf_bytes():
    messages = [
        FakeMessage("user", "Hello"),
        FakeMessage("assistant", "Hi there!"),
    ]

    file_bytes = build_pdf("Greeting", messages)

    assert file_bytes.startswith(b"%PDF")
    assert len(file_bytes) > 0


def test_build_pdf_handles_non_latin1_characters_without_crashing():
    messages = [
        FakeMessage("user", "📎 emoji and “curly quotes” shouldn't crash the PDF"),
    ]

    file_bytes = build_pdf("Unicode test", messages)

    assert file_bytes.startswith(b"%PDF")
