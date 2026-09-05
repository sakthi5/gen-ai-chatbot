import io

import pytest
from docx import Document as DocxDocument

from app.services.document_service import (
    extract_text,
    UnsupportedDocumentType,
    MAX_DOCUMENT_CHARACTERS,
)


def test_extract_text_from_txt():
    content = extract_text("notes.txt", b"Hello, this is a plain text file.")
    assert content == "Hello, this is a plain text file."


def test_extract_text_from_docx():
    buffer = io.BytesIO()
    document = DocxDocument()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(buffer)

    content = extract_text("report.docx", buffer.getvalue())

    assert "First paragraph." in content
    assert "Second paragraph." in content


def test_extract_text_from_pdf(monkeypatch):
    class FakePage:
        def extract_text(self):
            return "Text from a PDF page."

    class FakeReader:
        def __init__(self, _stream):
            self.pages = [FakePage()]

    monkeypatch.setattr("app.services.document_service.PdfReader", FakeReader)

    content = extract_text("paper.pdf", b"%PDF-fake-bytes")

    assert content == "Text from a PDF page."


def test_extract_text_rejects_unsupported_extension():
    with pytest.raises(UnsupportedDocumentType):
        extract_text("archive.zip", b"whatever")


def test_extract_text_truncates_long_documents():
    huge_text = "word " * (MAX_DOCUMENT_CHARACTERS // 4)  # well over the cap

    content = extract_text("big.txt", huge_text.encode("utf-8"))

    assert len(content) <= MAX_DOCUMENT_CHARACTERS + 200  # cap + truncation note
    assert content.endswith("...]")
